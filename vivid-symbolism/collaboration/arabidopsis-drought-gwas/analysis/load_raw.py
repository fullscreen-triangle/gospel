"""Extract the per-plant raw sheet and the easyGWAS tables to CSV for analysis."""
import sys
from pathlib import Path
import openpyxl
import pandas as pd

src = Path(sys.argv[1])
out = Path(sys.argv[2])
out.mkdir(parents=True, exist_ok=True)

wb = openpyxl.load_workbook(src, data_only=True, read_only=True)

ws = wb["Biomass data"]
rows = ws.iter_rows(values_only=True)
header = list(next(rows))
print("ALL BIOMASS HEADERS:")
for i, h in enumerate(header):
    print(f"  {i:2d} {h!r}")

raw, empty = [], 0
for r in rows:
    if r[0] is None and r[1] is None:
        empty += 1
        if empty > 50:
            break
        continue
    empty = 0
    raw.append(r[:8])
df = pd.DataFrame(raw, columns=header[:8])
df.to_csv(out / "raw_plants.csv", index=False)
print("\nraw rows:", len(df))
print(df.dtypes)
print("\nTreatment:", df["Treatment"].value_counts(dropna=False).to_dict())
print("Media:", df["Media"].value_counts(dropna=False).to_dict())
print("accessions:", df["Seed Line"].nunique(), "IDs:", df["ID"].nunique())
cell = df.groupby(["Seed Line", "Treatment", "Media"]).size()
print("plants per cell:", cell.describe().to_dict())
print("cells per accession:", df.groupby("Seed Line")[["Treatment", "Media"]]
      .apply(lambda g: len(g.drop_duplicates())).value_counts().to_dict())

for name in ["easyGWAS input file", "easyGWAS output", "Accessions", "Table 1", "Table 2"]:
    w = wb[name]
    it = w.iter_rows(values_only=True)
    h = list(next(it))
    data = [r for r in it if any(v is not None for v in r)]
    t = pd.DataFrame(data, columns=[str(x) for x in h])
    t.to_csv(out / (name.replace(" ", "_") + ".csv"), index=False)
    print(f"\n{name}: {t.shape}")
