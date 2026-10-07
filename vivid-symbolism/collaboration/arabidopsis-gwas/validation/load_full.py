"""Extract NS+DS.xlsx (per-plant data for all four cells) to a tidy CSV cache.

Columns of the tidy table: gid, line, treatment (Mock/WCS417), media (DS/NS),
br (replicate label, DS only), sheet, sheet_row, shoot, root, total.
"""
import openpyxl
import pandas as pd

from common import CACHE, SITE

FULL = SITE / "public" / "gwas" / "NS+DS.xlsx"


def extract_full(force=False):
    out = CACHE / "plants_all.csv"
    if out.exists() and not force:
        return pd.read_csv(out)
    CACHE.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(FULL, data_only=True, read_only=True)
    rows = []
    for name in wb.sheetnames:
        ws = wb[name]
        it = ws.iter_rows(values_only=True)
        hdr = [str(h).strip() if h is not None else "" for h in next(it)]
        col = {h: i for i, h in enumerate(hdr)}
        for k, r in enumerate(it, start=2):
            if r is None or all(v is None for v in r[:8]):
                continue
            g = lambda h: r[col[h]] if h in col and col[h] < len(r) else None
            rows.append({
                "sheet": name, "sheet_row": k, "gid": g("ID"), "line": g("Seed Line"),
                "treatment": g("Treatment"), "media": g("Media"), "br": g("BR"),
                "shoot": g("Shoot FW"), "root": g("Root FW"), "total": g("Total FW"),
            })
    df = pd.DataFrame(rows)
    for c in ["shoot", "root", "total"]:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    df["gid"] = pd.to_numeric(df["gid"], errors="coerce").astype("Int64")
    for c in ["treatment", "media", "br", "line"]:
        df[c] = df[c].astype("string").str.strip()
    df.to_csv(out, index=False)
    return df


if __name__ == "__main__":
    df = extract_full(force=True)
    pd.set_option("display.width", 200)
    print(df.shape)
    print(df.groupby(["sheet", "treatment", "media"], dropna=False).size())
    print("BR values:", df["br"].value_counts(dropna=False).to_dict())
    print("accessions:", df.groupby(["media", "treatment"])["gid"].nunique().to_dict())
    cell = df.groupby(["gid", "media", "treatment"]).size()
    print("plants per cell:", cell.describe().round(2).to_dict())
    print("cells per accession:", df.groupby("gid")[["media", "treatment"]].apply(lambda g: len(g.drop_duplicates())).value_counts().to_dict())
    print("missing weights:", df[["shoot", "root", "total"]].isna().sum().to_dict())
    print("BR x media x treatment:\n", df.groupby(["media", "treatment", "br"], dropna=False).size().unstack(fill_value=0))
    print("accessions per BR (DS):", df[df.media == "DS"].groupby("br")["gid"].nunique().to_dict())
    big = cell.sort_values(ascending=False).head(8)
    print("largest cells:\n", big)
