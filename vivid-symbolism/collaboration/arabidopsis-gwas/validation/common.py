"""Shared data loading for the WCS417 drought-rescue validation suite.

Every experiment reads the workbook through `load()` so that the audit, the
cell-mean recovery and the exclusion rules are applied identically everywhere.

Notation (per accession a, organ o):
    W   = mean fresh weight, WCS417-inoculated, drought        (raw plants available)
    M_D = mean fresh weight, mock, drought                     (recovered)
    M_N = mean fresh weight, mock, non-stress                  (recovered)
Recovered from the workbook's derived traits by the identities
    gain_D = W - M_D,   loss = 1 - M_D / M_N,
which are verified to rounding error in experiment E1.
"""
from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent                                   # collaboration/arabidopsis-gwas
SITE = ROOT.parents[1]                               # vivid-symbolism
WORKBOOK = SITE / "public" / "gwas" / "GWAS Kundai.xlsx"
CACHE = HERE / "_cache"
RESULTS = ROOT / "results"
FIGURES = ROOT / "figures"
SEED = 417                                           # the strain's name; fixed everywhere

ORGANS = {"Shoot": "SFW", "Root": "RFW", "Total": "TFW"}


def _sheet(wb, name):
    it = wb[name].iter_rows(values_only=True)
    header = [str(h) if h is not None else f"_c{i}" for i, h in enumerate(next(it))]
    rows = [r for r in it if any(v is not None for v in r)]
    return pd.DataFrame(rows, columns=header)


def extract(force: bool = False) -> None:
    """Workbook -> CSV cache (the raw sheet declares ~1e6 rows, so read once)."""
    CACHE.mkdir(parents=True, exist_ok=True)
    if (CACHE / "raw_plants.csv").exists() and not force:
        return
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    ws = wb["Biomass data"]
    rows = []
    for i, r in enumerate(ws.iter_rows(min_row=2, max_col=8, values_only=True), start=2):
        if any(v is not None for v in r):
            rows.append((i,) + tuple(r))
    raw = pd.DataFrame(rows, columns=["sheet_row", "ID", "line", "treatment", "media",
                                      "root_length", "shoot", "root", "total"])
    for c in ["root_length", "shoot", "root", "total"]:
        raw[c] = pd.to_numeric(raw[c], errors="coerce")
    raw.to_csv(CACHE / "raw_plants.csv", index=False)
    for name in ["easyGWAS input file", "easyGWAS output", "Accessions", "Table 1", "Table 2"]:
        _sheet(wb, name).to_csv(CACHE / (name.replace(" ", "_") + ".csv"), index=False)


@dataclass
class Data:
    raw: pd.DataFrame          # per-plant, WCS417 x drought only
    traits: pd.DataFrame       # workbook derived traits, one row per accession
    cells: pd.DataFrame        # recovered W, M_D, M_N per organ, indexed by genome id
    accessions: pd.DataFrame   # metadata (name, country, lat, lon), indexed by genome id
    candidates: pd.DataFrame   # easyGWAS output table
    table1: pd.DataFrame
    table2: pd.DataFrame
    repeated: list             # genome ids grown in >1 block (excluded from main set)
    main: list                 # genome ids in the analysis set


def load() -> Data:
    extract()
    raw = pd.read_csv(CACHE / "raw_plants.csv")
    tr = pd.read_csv(CACHE / "easyGWAS_input_file.csv")
    tr = tr.rename(columns={"IID": "gid"}).set_index("gid")
    cells = pd.DataFrame(index=tr.index)
    for o, a in ORGANS.items():
        W = tr[f"AVG_{o}_FW"].astype(float)
        MD = W - tr[f"AVG_Gain_in_{a}_Drought"].astype(float)
        MN = MD / (1.0 - tr[f"AVG_percent_Loss_in_{a}_under_drought"].astype(float))
        cells[f"{o}_W"], cells[f"{o}_MD"], cells[f"{o}_MN"] = W, MD, MN
    acc = pd.read_csv(CACHE / "Accessions.csv")
    acc["ID"] = pd.to_numeric(acc["ID"], errors="coerce")
    acc = acc.dropna(subset=["ID"]).drop_duplicates("ID")
    acc.index = acc["ID"].astype(int)
    cand = pd.read_csv(CACHE / "easyGWAS_output.csv")
    n_per = raw.groupby("ID").size()
    repeated = sorted(int(i) for i in n_per[n_per > 7].index)
    ok = cells[[f"{o}_{c}" for o in ["Shoot", "Root"] for c in ["W", "MD", "MN"]]].notna().all(axis=1)
    main = sorted(int(i) for i in cells.index[ok] if int(i) not in repeated)
    return Data(raw, tr, cells, acc, cand,
                pd.read_csv(CACHE / "Table_1.csv"), pd.read_csv(CACHE / "Table_2.csv"),
                repeated, main)


# --------------------------------------------------------------------------
# Response constructions. Each maps cell means to a number that is zero
# exactly when inoculation has no effect under drought (W = M_D).
# --------------------------------------------------------------------------
def gain(W, MD, MN):
    return W - MD


def increase(W, MD, MN):
    return W / MD - 1.0


def rescue(W, MD, MN):
    return (W - MD) / (MN - MD)


def f_lambda(W, MD, MN, lam):
    """Geometric interpolation: lam=1 -> increase, lam=0 -> rescue.
    Scale invariant for every lam (numerator and denominator both degree 1)."""
    return (W - MD) / (np.power(MD, lam) * np.power(MN - MD, 1.0 - lam))


CONSTRUCTIONS = {"gain": gain, "increase": increase, "rescue": rescue}


def save(name: str, obj) -> Path:
    RESULTS.mkdir(parents=True, exist_ok=True)
    p = RESULTS / f"{name}.json"

    def conv(o):
        if isinstance(o, (np.integer,)):
            return int(o)
        if isinstance(o, (np.floating,)):
            return None if not np.isfinite(o) else float(o)
        if isinstance(o, np.ndarray):
            return clean(o.tolist())
        if isinstance(o, float) and not np.isfinite(o):
            return None
        raise TypeError(type(o))

    def clean(o):
        if isinstance(o, dict):
            return {str(k): clean(v) for k, v in o.items()}
        if isinstance(o, (list, tuple)):
            return [clean(v) for v in o]
        if isinstance(o, float) and not np.isfinite(o):
            return None
        return o

    p.write_text(json.dumps(clean(obj), indent=2, default=conv), encoding="utf-8")
    return p
